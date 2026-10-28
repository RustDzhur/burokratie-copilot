#!/usr/bin/env python3
"""Convert the Markdown reports into .docx documents (tables, headings, lists).

Usage:
    md_to_docx.py docs/REPORT_2026-10-27.md ... --out reports/
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor


def add_inline(paragraph, text: str) -> None:
    """Handles **bold**, `code` and plain text without a full Markdown parser."""
    for token in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(token)


def convert(markdown: str, title: str) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    lines = markdown.splitlines()
    in_code = False
    code_buffer: list[str] = []
    table_buffer: list[str] = []

    def flush_table() -> None:
        nonlocal table_buffer
        if not table_buffer:
            return
        rows = [r.strip().strip("|").split("|") for r in table_buffer]
        rows = [r for r in rows if not all(set(c.strip()) <= {"-", ":", " "} for c in r)]
        if rows:
            width = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Light Grid Accent 1"
            for i, row in enumerate(rows):
                for j in range(width):
                    cell = table.cell(i, j)
                    cell.text = ""
                    add_inline(cell.paragraphs[0], (row[j] if j < len(row) else "").strip())
        table_buffer = []

    def flush_code() -> None:
        nonlocal code_buffer
        if code_buffer:
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_buffer))
            run.font.name = "Consolas"
            run.font.size = Pt(8)
        code_buffer = []

    for line in lines:
        if line.startswith("```"):
            if in_code:
                flush_code()
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_buffer.append(line)
            continue
        if line.strip().startswith("|"):
            table_buffer.append(line)
            continue
        flush_table()
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if level == 1:
                doc.add_heading(text, level=0)
            else:
                doc.add_heading(text, level=min(level, 4))
            continue
        if re.match(r"^[-*]\s+", stripped):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^[-*]\s+", "", stripped))
            continue
        if re.match(r"^\d+[.)]\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+[.)]\s+", "", stripped))
            continue
        p = doc.add_paragraph()
        add_inline(p, stripped)

    flush_table()
    flush_code()

    out = doc
    if title:
        out.core_properties.title = title
    return out


def main(argv: list[str]) -> int:
    args = argv[1:]
    out_dir = Path("reports")
    files: list[Path] = []
    i = 0
    while i < len(args):
        if args[i] == "--out":
            out_dir = Path(args[i + 1])
            i += 2
            continue
        files.append(Path(args[i]))
        i += 1
    out_dir.mkdir(parents=True, exist_ok=True)
    for src in files:
        if not src.exists():
            print(f"skip missing: {src}")
            continue
        doc = convert(src.read_text(encoding="utf-8"), src.stem)
        target = out_dir / f"{src.stem}.docx"
        doc.save(target)
        print(f"{src} -> {target}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
